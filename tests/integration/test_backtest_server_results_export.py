import datetime as dt
import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from infrastructure.config.settings import get_settings
from infrastructure.persistence.database import Base, create_session_factory


def _bar(day, index, open_, high, low, close, volume=1000.0):
    return {
        "date": dt.datetime.combine(day, dt.time(9, 15)) + dt.timedelta(minutes=index),
        "open": open_, "high": high, "low": low, "close": close, "volume": volume,
    }


def _write_sample_csv(path):
    day = dt.date(2026, 1, 5)
    rows = [_bar(day, i, 100.0, 100.0, 100.0, 100.0) for i in range(24)]
    rows.append(_bar(day, 24, 99.0, 99.0, 99.0, 99.0))
    rows.append(_bar(day, 25, 101.0, 101.0, 101.0, 101.0))
    rows.append(_bar(day, 26, 101.0, 103.5, 100.5, 103.0))
    lines = ["date,open,high,low,close,volume"]
    for r in rows:
        lines.append(
            f"{r['date'].isoformat()}+05:30,{r['open']},{r['high']},{r['low']},{r['close']},{r['volume']}"
        )
    path.write_text("\n".join(lines))


@pytest.fixture
def app(tmp_path, monkeypatch):
    db_path = tmp_path / "test.db"
    historical_dir = tmp_path / "historical_data"
    historical_dir.mkdir()
    _write_sample_csv(historical_dir / "TEST_historical.csv")
    _write_sample_csv(historical_dir / "TEST2_historical.csv")

    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{db_path}")
    monkeypatch.setenv("HISTORICAL_DATA_DIR", str(historical_dir))
    get_settings.cache_clear()

    import backtest_server

    Base.metadata.create_all(create_session_factory(f"sqlite:///{db_path}")().get_bind())
    yield backtest_server.create_app()

    get_settings.cache_clear()


RUN_PAYLOAD = {
    "strategy": "test_json_threshold",
    "symbols": ["TEST", "TEST2"],
    "quantity": 50,
    "date_from": "2026-01-01",
    "date_to": "2026-01-10",
}


def test_export_selected_only_includes_the_given_ids(app):
    with TestClient(app) as client:
        run_res = client.post("/api/backtest/run", json=RUN_PAYLOAD).json()
        test_id = run_res["saved_result_ids"]["TEST"]

        res = client.post("/api/backtest/results/export-selected", json={"result_ids": [test_id]})

        assert res.status_code == 200
        assert res.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        wb = openpyxl.load_workbook(filename=io.BytesIO(res.content))
        sheet = wb.worksheets[0]
        symbol_col = [c.value for c in sheet[1]].index("Symbols") + 1
        data_rows = list(sheet.iter_rows(min_row=2, values_only=False))
        assert len(data_rows) == 1
        assert data_rows[0][symbol_col - 1].value == "TEST"


def test_export_selected_with_both_symbols(app):
    with TestClient(app) as client:
        run_res = client.post("/api/backtest/run", json=RUN_PAYLOAD).json()
        ids = list(run_res["saved_result_ids"].values())

        res = client.post("/api/backtest/results/export-selected", json={"result_ids": ids})

    assert res.status_code == 200
    wb = openpyxl.load_workbook(filename=io.BytesIO(res.content))
    sheet = wb.worksheets[0]
    assert sheet.max_row - 1 == 2  # header + 2 data rows (TEST, TEST2)


def test_export_selected_rejects_empty_id_list(app):
    with TestClient(app) as client:
        res = client.post("/api/backtest/results/export-selected", json={"result_ids": []})

    assert res.status_code == 422  # min_length=1 on the request model


def test_export_selected_with_all_unknown_ids_returns_400(app):
    with TestClient(app) as client:
        res = client.post("/api/backtest/results/export-selected", json={"result_ids": [999999]})

    assert res.status_code == 400
