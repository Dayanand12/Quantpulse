import datetime as dt
import json
from io import BytesIO

import openpyxl
import pandas as pd

from core.domain.backtest_result import BacktestResult, BacktestRunParams
from runners.backtesting.report import export_stored_results_to_excel


def _result(id_, params_json, metrics, created_at=None):
    run_params = BacktestRunParams(
        strategy_name="vwap_reclaim",
        symbols="RELIANCE,TCS",
        timeframe="minute",
        date_from=dt.date(2026, 1, 1),
        date_to=dt.date(2026, 6, 30),
        quantity=50,
        stoploss_pct=0.8,
        target_pct=2.0,
        trailing_pct=1.0,
        max_cycles_per_day=10,
        start_time="09:20",
        end_time="11:30",
        charges_enabled=True,
        strategy_params_json=json.dumps(params_json),
    )
    return BacktestResult(
        params=run_params,
        result={"metrics": metrics},
        id=id_,
        created_at=created_at or dt.datetime(2026, 7, 1, 10, 0),
    )


def _read_sheet(content: bytes) -> pd.DataFrame:
    return pd.read_excel(BytesIO(content), sheet_name="Parameter Comparison")


def test_exports_one_row_per_stored_result():
    results = [
        _result(1, {"parameters": {"adx_threshold": 15}}, {"total_trades": 20, "profit_factor": 1.5}),
        _result(2, {"parameters": {"adx_threshold": 25}}, {"total_trades": 10, "profit_factor": 1.2}),
    ]

    df = _read_sheet(export_stored_results_to_excel(results))

    assert len(df) == 2
    assert set(df["Result ID"]) == {1, 2}
    assert set(df["adx_threshold"]) == {15, 25}
    assert set(df["Total Trades"]) == {20, 10}


def test_union_of_parameter_keys_across_results_with_blanks_for_missing():
    results = [
        _result(1, {"parameters": {"adx_threshold": 15}}, {"total_trades": 5}),
        _result(2, {"parameters": {"volume_ratio_min": 2.0}}, {"total_trades": 7}),
    ]

    df = _read_sheet(export_stored_results_to_excel(results))

    assert {"adx_threshold", "volume_ratio_min"} <= set(df.columns)
    row1 = df[df["Result ID"] == 1].iloc[0]
    row2 = df[df["Result ID"] == 2].iloc[0]
    assert row1["adx_threshold"] == 15
    assert pd.isna(row1["volume_ratio_min"])
    assert row2["volume_ratio_min"] == 2.0
    assert pd.isna(row2["adx_threshold"])


def test_handles_result_with_no_strategy_params_json():
    results = [_result(1, {}, {"total_trades": 3})]
    results[0] = BacktestResult(
        params=BacktestRunParams(
            strategy_name="orb_reversal", symbols="WATCHLIST", timeframe="minute",
            date_from=dt.date(2026, 1, 1), date_to=dt.date(2026, 6, 30), quantity=50,
            stoploss_pct=0.8, target_pct=2.0, trailing_pct=1.0, max_cycles_per_day=10,
            start_time="09:20", end_time="11:30", charges_enabled=True, strategy_params_json="",
        ),
        result={"metrics": {"total_trades": 3}},
        id=1,
    )

    df = _read_sheet(export_stored_results_to_excel(results))

    assert len(df) == 1
    assert df.iloc[0]["Total Trades"] == 3


def test_includes_run_settings_and_metrics_columns():
    results = [_result(1, {"parameters": {}}, {
        "total_trades": 10, "win_rate": 55.0, "profit_factor": 1.3, "total_pnl": 5000.0,
        "sharpe_ratio": 1.2, "max_drawdown_pct": 3.1,
    })]

    df = _read_sheet(export_stored_results_to_excel(results))
    row = df.iloc[0]

    assert row["Symbols"] == "RELIANCE,TCS"
    assert row["Stop Loss %"] == 0.8
    assert row["Target %"] == 2.0
    assert row["Win Rate %"] == 55.0
    assert row["Sharpe Ratio"] == 1.2
    assert row["Max Drawdown %"] == 3.1


def test_column_guide_sheet_categorizes_every_column():
    results = [_result(1, {"parameters": {"adx_threshold": 15}}, {"total_trades": 20, "profit_factor": 1.5})]
    content = export_stored_results_to_excel(results)

    guide = pd.read_excel(BytesIO(content), sheet_name="Column Guide")
    by_column = dict(zip(guide["Column"], guide["Type"]))

    assert by_column["Result ID"] == "Identifier"
    assert by_column["Created At"] == "Identifier"
    assert by_column["Stop Loss %"] == "Input"
    assert by_column["Date From"] == "Input"
    assert by_column["adx_threshold"] == "Input"  # dynamic strategy parameter -> still an input
    assert by_column["Total Trades"] == "Output"
    assert by_column["Sharpe Ratio"] == "Output"


def test_parameter_comparison_sheet_stays_the_default_first_sheet():
    # parse_scenarios_from_excel() reads pd.read_excel's default
    # sheet_name=0 — the Column Guide sheet must never end up first.
    results = [_result(1, {"parameters": {}}, {"total_trades": 1})]
    content = export_stored_results_to_excel(results)

    wb = openpyxl.load_workbook(BytesIO(content))
    assert wb.sheetnames[0] == "Parameter Comparison"
    assert "Column Guide" in wb.sheetnames


def test_columns_are_color_coded_by_category():
    results = [_result(1, {"parameters": {"adx_threshold": 15}}, {"total_trades": 20})]
    content = export_stored_results_to_excel(results)

    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb["Parameter Comparison"]
    header = {cell.value: cell.column for cell in ws[1]}

    result_id_col = header["Result ID"]
    stoploss_col = header["Stop Loss %"]
    param_col = header["adx_threshold"]
    trades_col = header["Total Trades"]

    # Header row
    assert ws.cell(row=1, column=result_id_col).fill.start_color.rgb == "00F2F2F2"
    assert ws.cell(row=1, column=stoploss_col).fill.start_color.rgb == "00DCE6F1"
    assert ws.cell(row=1, column=param_col).fill.start_color.rgb == "00DCE6F1"
    assert ws.cell(row=1, column=trades_col).fill.start_color.rgb == "00E2EFDA"

    # Data row (row 2) — colored too, not just the header, so scrolling
    # past the header still shows which group a cell belongs to.
    assert ws.cell(row=2, column=result_id_col).fill.start_color.rgb == "00F2F2F2"
    assert ws.cell(row=2, column=stoploss_col).fill.start_color.rgb == "00DCE6F1"
    assert ws.cell(row=2, column=trades_col).fill.start_color.rgb == "00E2EFDA"
